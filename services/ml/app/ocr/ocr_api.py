import os, io, time, shutil
import tempfile, threading, contextlib
import numpy as np
import torch
from PIL import Image
from typing import Generator, List, Tuple
from fastapi import FastAPI, File, UploadFile, HTTPException
from pptx import Presentation
from doctr.io import DocumentFile
from doctr.models import ocr_predictor
from huggingface_hub import login

from .ocr_config import (
    MODEL_DIR,
    HF_TOKEN,
    DEVICE,
    OCR_DET_ARCH,
    OCR_RECO_ARCH,
    OCR_CONCURRENCY_LIMIT,
    OCR_MAX_PAGES,
    OCR_ASSUME_STRAIGHT,
    TORCH_NUM_THREADS,
    TORCH_NUM_INTEROP,
    OCR_PDF_DPI,
    OCR_BATCH_PAGES,
    SEM_TIMEOUT_S,
    MAX_UPLOAD_MB,
    PPTX_MAX_SLIDES,
)


# -------------- Device Config --------------
def _select_device(device_str: str) -> str:
    """
    If it builds on CUDA and no GPU is detected, then it will remain on the CPU.
    """
    device = device_str or "cpu"
    if device.startswith("cuda") and not torch.cuda.is_available():
        print("[OCR boot] CUDA requested but not available, falling back to CPU.")
        return "cpu"
    return device

device = _select_device(DEVICE)

# FP16 effective: only if you are actually on CUDA
EFFECTIVE_FP16 = device.startswith("cuda") and torch.cuda.is_available()


# -------------- Global Config --------------
# Configure PyTorch threading based on configuration values.
# If the host doesn't support these options, ignore the error and continue.
try:
    torch.set_num_threads(TORCH_NUM_THREADS)
    torch.set_num_interop_threads(TORCH_NUM_INTEROP)
except Exception:
    # Do not fail the application startup if thread configuration is not supported.
    pass

# Optional Hugging Face auth: Not required for public docTR weights, 
# but useful for private repositories.
if HF_TOKEN:
    try:
        login(token=HF_TOKEN)
    except Exception as e:
        print(f"[OCR boot] HF login failed (non-fatal): {e}")

# Boot-time summary for observability and debugging.
print(
    "[OCR boot] "
    f"device={device} | det={OCR_DET_ARCH} | reco={OCR_RECO_ARCH} | "
    f"fp16={EFFECTIVE_FP16} | straight={OCR_ASSUME_STRAIGHT} | "
    f"max_pages={OCR_MAX_PAGES} | batch_pages={OCR_BATCH_PAGES} | "
    f"concurrency={OCR_CONCURRENCY_LIMIT} | pdf_dpi={OCR_PDF_DPI}"
)


# -------------- Build docTR predictor --------------
def _build_predictor(to_device):
    """
    Build a docTR OCR predictor and move it to the target device.

    assume_straight_pages=True skips expensive geometric corrections, which
    improves performance on documents that are mostly straight.
    """
    mdl = ocr_predictor(
        det_arch=OCR_DET_ARCH,
        reco_arch=OCR_RECO_ARCH,
        pretrained=True,
        assume_straight_pages=OCR_ASSUME_STRAIGHT,
    ).to(to_device)
    mdl.eval()  # inference-only mode
    return mdl

# Main OCR model instantiated at import time.
model = _build_predictor(device)

def _get_autocast_context(device, use_fp16: bool):
    """
    Return an autocast context manager.
    - On GPU + use_fp16=True => torch.autocast(float16)
    - Otherwise => nullcontext (no autocast)
    """
    is_cuda = False
    if isinstance(device, str):
        is_cuda = device.startswith("cuda")
    elif hasattr(device, "type"):
        is_cuda = (device.type == "cuda")

    if use_fp16 and is_cuda:
        return torch.autocast(device_type="cuda", dtype=torch.float16)

    return contextlib.nullcontext()


# -------------- Concurrency guard (per worker) --------------
# Concurrency limiter: allows only OCR_CONCURRENCY_LIMIT OCR requests to run at the same time
SEM = threading.BoundedSemaphore(OCR_CONCURRENCY_LIMIT)

class AcquireWithTimeout:
    """
    Context manager to acquire a semaphore with a timeout.

    If the semaphore cannot be acquired within timeout_s seconds, it raises
    an HTTP 429 error to signal too many concurrent requests.
    """

    def __init__(self, sem: threading.BoundedSemaphore, timeout_s: int):
        self.sem = sem
        self.timeout_s = timeout_s
        self.acquired = False

    def __enter__(self):
        # Return 429 if the semaphore cannot be acquired in time.
        self.acquired = self.sem.acquire(timeout=self.timeout_s)
        if not self.acquired:
            raise HTTPException(
                status_code=429,
                detail="Too many concurrent requests, please retry.",
            )
        return self

    def __exit__(self, exc_type, exc, tb):
        if self.acquired:
            self.sem.release()
        # Do not suppress exceptions.
        return False


# -------------- Helpers --------------
def _reject_if_too_big(path: str):
    """
    Enforce an upper limit on uploaded file size.

    This prevents unnecessary decoding and reduces the risk of OOM errors on
    very large documents.
    """
    size_mb = os.path.getsize(path) / (1024 * 1024)
    if size_mb > MAX_UPLOAD_MB:
        raise HTTPException(
            status_code=413,
            detail=f"File too large ({size_mb:.1f} MB > {MAX_UPLOAD_MB} MB)",
        )
    
def _load_pdf_document(path: str) -> DocumentFile:
    """
    Load a PDF into a DocumentFile.

    Prefer the 'scale' argument (pypdfium2 backend). If the installed version
    doesn't support it, fall back to 'dpi' for older backends.
    """
    scale = OCR_PDF_DPI / 72.0  # pypdfium2: scale = dpi / 72
    try:
        return DocumentFile.from_pdf(path, scale=scale)
    except TypeError:
        return DocumentFile.from_pdf(path, dpi=OCR_PDF_DPI)

def _extract_text_from_pptx(file_path: str) -> str:
    """
    Fast path for PPTX files: extract slide text directly without OCR.

    This is significantly faster and more robust than rasterizing slides,
    as long as text is stored as PPTX shapes.
    """
    prs = Presentation(file_path)
    chunks: List[str] = []
    for i, slide in enumerate(prs.slides):
        if i >= PPTX_MAX_SLIDES:
            break
        for shape in slide.shapes:
            if hasattr(shape, "text") and shape.text:
                chunks.append(shape.text)
    return "\n".join(chunks)

def _iter_doc_chunks(doc: DocumentFile, chunk_size: int) -> Generator[DocumentFile, None, None]:
    """
    Yield DocumentFile slices in fixed-size page chunks.

    This allows us to process large documents in smaller batches, controlling
    memory usage and improving latency characteristics.
    """
    n = len(doc)
    for start in range(0, n, chunk_size):
        yield doc[start : min(start + chunk_size, n)]

def _pages_to_text(result) -> str:
    """
    Flatten docTR result structure (page → block → line → word) into plain text.
    Words are joined by newlines to avoid losing all structure while staying simple.
    """
    words: List[str] = []
    for page in result.pages:
        for block in page.blocks:
            for line in block.lines:
                for w in line.words:
                    if getattr(w, "value", None):
                        words.append(w.value)
    return "\n".join(words)

def _predict_text_with_model(
    mdl: torch.nn.Module,
    device: str,
    doc: DocumentFile,
    use_fp16: bool,
    batch_pages: int,
) -> str:
    """
    Run OCR on a DocumentFile with page batching.
    - Truncates the document to OCR_MAX_PAGES to keep runtime bounded.
    - Processes pages in chunks of 'batch_pages' to control memory usage.
    Returns a newline-separated text string.
    """
    if len(doc) > OCR_MAX_PAGES:
        doc = doc[:OCR_MAX_PAGES]

    out_chunks: List[str] = []
    with torch.inference_mode():
        for subdoc in _iter_doc_chunks(doc, batch_pages):
            with _get_autocast_context(device, use_fp16):
                result = mdl(subdoc)
            out_chunks.append(_pages_to_text(result))
    return "\n".join(filter(None, out_chunks))

def _run_ocr(doc: DocumentFile) -> Tuple[str, str]:
    """
    Execute OCR and return (text, device_used).
    """
    text = _predict_text_with_model(
        mdl=model,
        device=device,
        doc=doc,
        use_fp16=EFFECTIVE_FP16,
        batch_pages=OCR_BATCH_PAGES,
    )
    return text, device


# -------------- FastAPI application --------------
app = FastAPI(title="📝 OCR API", version="1.0.0")

@app.on_event("startup")
async def startup():
    """
    Best-effort warmup of the OCR model.

    Warmup runs a dummy inference to populate caches and reduce latency on the
    first real request. Any failure is logged but not treated as fatal.
    """
    try:
        try:
            # Primary path: generate a white test image in memory.
            img = Image.fromarray((255 * np.ones((64, 256, 3), dtype="uint8")))
            bio = io.BytesIO()
            img.save(bio, format="PNG")
            dummy_doc = DocumentFile.from_images([bio.getvalue()])
        except Exception:
            # Fallback path: minimal fake PNG buffer if PIL / numpy creation fails.
            bio = io.BytesIO(b"\x89PNG\r\n\x1a\n" + b"\x00" * 64)
            dummy_doc = DocumentFile.from_images([bio])

        with torch.inference_mode(), _get_autocast_context(device, EFFECTIVE_FP16):
            _ = model(dummy_doc)

        print(f"[---OCR---] Warmup completed. Device = {device}")
    except Exception as e:
        print(f"[---OCR---] Warmup failed (non-fatal): {e}")

@app.get("/healthz")
def healthz():
    """
    Lightweight health endpoint exposing runtime and OCR configuration.
    """
    return {
        "ok": True,
        "model_cache_dir": MODEL_DIR,
        "device": device,
        "det_arch": OCR_DET_ARCH,
        "reco_arch": OCR_RECO_ARCH,
        "fp16": EFFECTIVE_FP16,
        "assume_straight": OCR_ASSUME_STRAIGHT,
        "max_pages": OCR_MAX_PAGES,
        "batch_pages": OCR_BATCH_PAGES,
        "concurrency": OCR_CONCURRENCY_LIMIT,
        "pdf_dpi": OCR_PDF_DPI,
        "torch_num_threads": torch.get_num_threads(),
        "torch_num_interop_threads": torch.get_num_interop_threads()
        if hasattr(torch, "get_num_interop_threads")
        else None,
    }

@app.post("/")
async def extract_text(file: UploadFile = File(...)):
    """
    OCR extraction endpoint.
    Accepts:
        - PDF
        - JPEG / JPG / PNG
        - PPTX (fast text extraction without OCR)
    Returns:
        - extracted_text
        - timings (total / load / inference in ms)
        - basic metadata (filename, extension, pages, device, model config)
    """
    filename = file.filename or ""
    ext = filename.rsplit(".", 1)[-1].lower() if "." in filename else ""
    supported = {"pdf", "jpeg", "jpg", "png", "pptx"}
    if ext not in supported:
        raise HTTPException(status_code=400, detail=f"Unsupported file format: .{ext}")

    # Save to a temporary file for libraries that require a filesystem path.
    with tempfile.NamedTemporaryFile(delete=False, suffix=f".{ext}") as tmp:
        shutil.copyfileobj(file.file, tmp)
        temp_path = tmp.name

    try:
        _reject_if_too_big(temp_path)

        # Concurrency guard with timeout -> returns 429 if saturated.
        with AcquireWithTimeout(SEM, SEM_TIMEOUT_S):
            t0 = time.monotonic()

            # -------- Rasterization / loading phase --------
            t_load0 = time.monotonic()
            pages = 0

            if ext in {"jpeg", "jpg", "png"}:
                # Image input: load as DocumentFile from image path.
                doc = DocumentFile.from_images(temp_path)
                pages = len(doc)

            elif ext == "pdf":
                # PDF input: load pages via DocumentFile with proper DPI/scale.
                doc = _load_pdf_document(temp_path)
                pages = len(doc)

            elif ext == "pptx":
                # Fast path for PPTX: native text extraction, no OCR call.
                text = _extract_text_from_pptx(temp_path)
                ms_total = int((time.monotonic() - t0) * 1000)
                return {
                    "filename": filename,
                    "ext": ext,
                    "pages": None,
                    "device_used": None,
                    "extracted_text": text,
                    "timings_ms": {
                        "total": ms_total,
                        "raster_or_load": int((time.monotonic() - t_load0) * 1000),
                        "inference": 0,
                    },
                }

            else:
                # Defensive: should never happen due to earlier extension check.
                raise HTTPException(status_code=400, detail="Unsupported file type")

            t_load_ms = int((time.monotonic() - t_load0) * 1000)

            # -------- Inference phase --------
            t_inf0 = time.monotonic()
            text, device_used = _run_ocr(doc)
            t_inf_ms = int((time.monotonic() - t_inf0) * 1000)

            ms_total = int((time.monotonic() - t0) * 1000)
            return {
                "filename": filename,
                "ext": ext,
                "pages": pages,
                "device_used": device_used,
                "det_arch": OCR_DET_ARCH,
                "reco_arch": OCR_RECO_ARCH,
                "fp16": EFFECTIVE_FP16,
                "extracted_text": text,
                "timings_ms": {
                    "total": ms_total,
                    "raster_or_load": t_load_ms,
                    "inference": t_inf_ms,
                },
            }

    finally:
        # Always clean up the temporary file, regardless of success or failure.
        try:
            os.remove(temp_path)
        except Exception:
            # Ignore cleanup errors.
            pass
