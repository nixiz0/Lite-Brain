from pptx import Presentation

def extract_pptx_pages(path: str):
    """
    Extracts text from a PPTX file slide-by-slide in order.
    Returns a list of dicts:
        [{"page_number": n, "text": "..."}]
    """
    prs = Presentation(path)
    pages = []

    for idx, slide in enumerate(prs.slides, start=1):
        texts = []

        # Collect text from all shapes that contain textual content
        for shape in slide.shapes:
            if hasattr(shape, "text") and shape.text:
                t = shape.text.strip()
                if t:
                    texts.append(t)

        # Merge all extracted blocks into a single text string
        page_text = "\n".join(texts).strip()

        # Preserve slide numbering even if slide contains no text
        pages.append(
            {
                "page_number": idx,
                "text": page_text,
            }
        )

    return pages
